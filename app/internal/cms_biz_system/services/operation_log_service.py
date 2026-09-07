"""
操作日志服务

对外暴露：
    write_log()   — 写入一条日志（供各路由埋点调用，内部捕获异常，不影响主业务）
    list_logs()   — 分页查询日志列表
    get_log()     — 获取单条日志详情
    clear_logs()  — 按时间范围删除日志
    export_logs_excel() — 导出 Excel（bytes）

操作类型常量集中定义在 OperationType 中，方便未来路由埋点直接引用。

读时翻译（read-time i18n）：
    encode_log_i18n(code, **params) → 构造 "@i18n:..." sentinel 值入库
    translate_log_value(value)       → 读时解析 sentinel，按请求语言翻译
    向后兼容：无 sentinel 前缀的旧数据原样直通。
"""

import io
import json as _json
import re
from datetime import datetime

from loguru import logger
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from ..models.operation_log import OperationLog
from ..models.user import User
from app.internal.cms_biz_system.schemas.operation_log import OperationLogItem
from app.common.schemas import PaginatedResponse
from app.common.core.exceptions import ErrorCode, NotFoundException
from app.common.core.i18n import get_msg, get_message
from app.config import app_tz

# ── 读时翻译 sentinel ──────────────────────────────────────
LOG_I18N_PREFIX = "@i18n:"


def encode_log_i18n(code: str, **params) -> str:
    """将 i18n code + 参数编码为 sentinel 字符串，写入 operation_object / operation_content / error_message 字段。"""
    return LOG_I18N_PREFIX + _json.dumps({"code": code, "params": params}, ensure_ascii=False)


def translate_log_value(value: str | None) -> str | None:
    """读时翻译：若值以 LOG_I18N_PREFIX 开头则解析并根据请求语言翻译，否则原样返回。"""
    if not value or not value.startswith(LOG_I18N_PREFIX):
        return value
    try:
        payload_str = value[len(LOG_I18N_PREFIX):]
        payload = _json.loads(payload_str)
        return get_msg(payload["code"], **payload.get("params", {}))
    except Exception:
        return value


def translate_legacy_log_value(value: str | None, lang: str = "cn") -> str | None:
    """Excel 导出兜底：翻译历史遗留的裸 log.* 编码（含冒号参数变体）。

    仅用于导出场景——列表/详情 API 的裸编码由前端翻译，此处不介入。
    规则与前端 translateLogContent 保持一致：
    - log.movie.*:类型码 → 带类型标签的扁平 key
    - log.field.edit / log.i18n.edit:语言 → "编辑（语言）"
    - log.episode.* / log.review.reject:参数 → "文案:参数"
    - 无翻译时原样返回
    """
    if not value or not value.startswith("log."):
        return value
    colon = value.find(":")
    if colon == -1:
        translated = get_message(value, lang=lang)
        return translated
    base, param = value[:colon], value[colon + 1:]
    if base in ("log.movie.create", "log.movie.edit", "log.movie.delete"):
        typed_key = f"{base}.{param}"
        typed = get_message(typed_key, lang=lang)
        if typed != typed_key:
            return typed
        return value
    if base in ("log.field.edit", "log.i18n.edit"):
        return get_message(f"{base}.lang", lang=lang, lang_code=param)
    if base in ("log.episode.inject", "log.episode.remove"):
        return get_message(base, lang=lang, title=param)
    if base == "log.review.reject":
        return get_message(base, lang=lang, reason=param)
    base_translated = get_message(base, lang=lang)
    if base_translated == base:
        return value
    return f"{base_translated}:{param}"


# 历史遗留裸 LOG_* 大写编码（早期 cms_biz_scp 直接写入 operation_content 的值）
_LEGACY_UPPER_CODE_RE = re.compile(r"^LOG_[A-Z0-9_]+$")


def _translate_for_export(value: str | None, lang: str) -> str:
    """导出用组合翻译：sentinel → 历史遗留裸 log.* → 历史遗留裸 LOG_* 大写编码。

    列表/详情 API 的裸编码由前端翻译（读时仅翻译 sentinel），
    Excel 导出无前端介入，需在此按导出语言组合兜底翻译存量数据。
    无匹配翻译时原样返回。
    """
    if not value:
        return ""
    translated = translate_log_value(value)
    if translated != value:
        return translated
    translated = translate_legacy_log_value(value, lang)
    if translated != value:
        return translated
    if _LEGACY_UPPER_CODE_RE.match(value):
        upper_translated = get_message(value, lang=lang)
        if upper_translated != value:
            return upper_translated
    return value

# ── 操作类型常量 ────────────────────────────────────────────
class OperationType:
    USER_LOGIN = "USER_LOGIN"
    USER_LOGOUT = "USER_LOGOUT"
    USER_CHANGE_PWD = "USER_CHANGE_PWD"

    USER_CREATE = "USER_CREATE"
    USER_EDIT = "USER_EDIT"
    USER_DELETE = "USER_DELETE"
    USER_RESET_PWD = "USER_RESET_PWD"
    USER_STATUS = "USER_STATUS"

    ROLE_CREATE = "ROLE_CREATE"
    ROLE_EDIT = "ROLE_EDIT"
    ROLE_DELETE = "ROLE_DELETE"
    ROLE_PERM = "ROLE_PERM"
    ROLE_STATUS = "ROLE_STATUS"

    DICT_CREATE = "DICT_CREATE"
    DICT_EDIT = "DICT_EDIT"
    DICT_DELETE = "DICT_DELETE"
    DICT_STATUS = "DICT_STATUS"

    CONFIG_CREATE = "CONFIG_CREATE"
    CONFIG_EDIT = "CONFIG_EDIT"
    CONFIG_DELETE = "CONFIG_DELETE"

    CONTENT_CREATE = "CONTENT_CREATE"
    CONTENT_EDIT = "CONTENT_EDIT"
    CONTENT_DELETE = "CONTENT_DELETE"
    CONTENT_PUBLISH = "CONTENT_PUBLISH"
    CONTENT_UNPUBLISH = "CONTENT_UNPUBLISH"
    CONTENT_REVIEW_INITIATE = "CONTENT_REVIEW_INITIATE"
    CONTENT_REVIEW_APPROVE = "CONTENT_REVIEW_APPROVE"
    CONTENT_REVIEW_REJECT = "CONTENT_REVIEW_REJECT"

    CHANNEL_UPDATE = "CHANNEL_UPDATE"

    PHYSICAL_CHANNEL_CREATE = "PHYSICAL_CHANNEL_CREATE"
    PHYSICAL_CHANNEL_DELETE = "PHYSICAL_CHANNEL_DELETE"

    CHANNEL_METADATA_CREATE = "CHANNEL_METADATA_CREATE"
    CHANNEL_METADATA_UPDATE = "CHANNEL_METADATA_UPDATE"
    CHANNEL_METADATA_DELETE = "CHANNEL_METADATA_DELETE"

    SCHEDULE_METADATA_CREATE = "SCHEDULE_METADATA_CREATE"
    SCHEDULE_METADATA_UPDATE = "SCHEDULE_METADATA_UPDATE"
    SCHEDULE_METADATA_DELETE = "SCHEDULE_METADATA_DELETE"

    CAST_ROLE_MAP_CREATE = "CAST_ROLE_MAP_CREATE"
    CAST_ROLE_MAP_UPDATE = "CAST_ROLE_MAP_UPDATE"
    CAST_ROLE_MAP_DELETE = "CAST_ROLE_MAP_DELETE"
    CAST_ROLE_MAP_LINK = "CAST_ROLE_MAP_LINK"
    CAST_ROLE_MAP_UNLINK = "CAST_ROLE_MAP_UNLINK"

    CONTENT_PACKAGE_LINK = "CONTENT_PACKAGE_LINK"
    CONTENT_PACKAGE_UNLINK = "CONTENT_PACKAGE_UNLINK"
    CONTENT_CATEGORY_LINK = "CONTENT_CATEGORY_LINK"
    CONTENT_CATEGORY_UNLINK = "CONTENT_CATEGORY_UNLINK"

    POSTER_UPLOAD = "POSTER_UPLOAD"
    POSTER_DELETE = "POSTER_DELETE"
    PICTURE_PUBLISH = "PICTURE_PUBLISH"

    CHANNEL_FIELD_UPDATE = "CHANNEL_FIELD_UPDATE"
    CHANNEL_I18N_UPDATE = "CHANNEL_I18N_UPDATE"

    SCHEDULE_FIELD_UPDATE = "SCHEDULE_FIELD_UPDATE"
    SCHEDULE_I18N_UPDATE = "SCHEDULE_I18N_UPDATE"

    PROGRAM_METADATA_CREATE = "PROGRAM_METADATA_CREATE"
    PROGRAM_METADATA_UPDATE = "PROGRAM_METADATA_UPDATE"
    PROGRAM_METADATA_DELETE = "PROGRAM_METADATA_DELETE"

    SERIES_METADATA_CREATE = "SERIES_METADATA_CREATE"
    SERIES_METADATA_UPDATE = "SERIES_METADATA_UPDATE"
    SERIES_METADATA_DELETE = "SERIES_METADATA_DELETE"

    MOVIE_CREATE = "MOVIE_CREATE"
    MOVIE_UPDATE = "MOVIE_UPDATE"
    MOVIE_DELETE = "MOVIE_DELETE"

    VOD_FIELD_UPDATE = "VOD_FIELD_UPDATE"
    VOD_I18N_UPDATE = "VOD_I18N_UPDATE"

    MOVIE_FIELD_UPDATE = "MOVIE_FIELD_UPDATE"
    MOVIE_I18N_UPDATE = "MOVIE_I18N_UPDATE"

    CAST_ROLE_MAP_FIELD_UPDATE = "CAST_ROLE_MAP_FIELD_UPDATE"
    CAST_ROLE_MAP_I18N_UPDATE = "CAST_ROLE_MAP_I18N_UPDATE"

    EPISODE_INJECT = "EPISODE_INJECT"
    EPISODE_REMOVE = "EPISODE_REMOVE"

    SEASON_SERIES_INJECT = "SEASON_SERIES_INJECT"
    SEASON_SERIES_REMOVE = "SEASON_SERIES_REMOVE"

    PROVIDER_CREATE = "PROVIDER_CREATE"
    PROVIDER_EDIT = "PROVIDER_EDIT"
    PROVIDER_DELETE = "PROVIDER_DELETE"

    CONTRACT_CREATE = "CONTRACT_CREATE"
    CONTRACT_EDIT = "CONTRACT_EDIT"
    CONTRACT_DELETE = "CONTRACT_DELETE"

    LICENSE_CREATE = "LICENSE_CREATE"
    LICENSE_EDIT = "LICENSE_EDIT"
    LICENSE_DELETE = "LICENSE_DELETE"

    PACKAGE_CREATE = "PACKAGE_CREATE"
    PACKAGE_EDIT = "PACKAGE_EDIT"
    PACKAGE_DELETE = "PACKAGE_DELETE"
    PACKAGE_CONTENT_ADD = "PACKAGE_CONTENT_ADD"
    PACKAGE_CONTENT_REMOVE = "PACKAGE_CONTENT_REMOVE"
    PACKAGE_EXPORT = "PACKAGE_EXPORT"
    PACKAGE_IMPORT = "PACKAGE_IMPORT"
    PACKAGE_TEMPLATE_DOWNLOAD = "PACKAGE_TEMPLATE_DOWNLOAD"

    PUBLISH_NOW = "PUBLISH_NOW"
    UNPUBLISH_NOW = "UNPUBLISH_NOW"
    PUBLISH_BATCH = "PUBLISH_BATCH"
    UNPUBLISH_BATCH = "UNPUBLISH_BATCH"
    PUBLISH_PLAN_CREATE = "PUBLISH_PLAN_CREATE"
    PUBLISH_PLAN_UPDATE = "PUBLISH_PLAN_UPDATE"
    PUBLISH_PLAN_CANCEL = "PUBLISH_PLAN_CANCEL"
    PUBLISH_PLAN_EXECUTE = "PUBLISH_PLAN_EXECUTE"
    UNPUBLISH_PLAN_EXECUTE = "UNPUBLISH_PLAN_EXECUTE"

    SCHEDULE_CREATE = "SCHEDULE_CREATE"
    SCHEDULE_UPDATE = "SCHEDULE_UPDATE"
    SCHEDULE_ARCHIVE = "SCHEDULE_ARCHIVE"
    SCHEDULE_DELETE = "SCHEDULE_DELETE"
    SCHEDULE_EXPORT = "SCHEDULE_EXPORT"
    SCHEDULE_IMPORT = "SCHEDULE_IMPORT"

    CAST_CREATE = "CAST_CREATE"
    CAST_EDIT = "CAST_EDIT"
    CAST_DELETE = "CAST_DELETE"

    CATEGORY_CREATE = "CATEGORY_CREATE"
    CATEGORY_EDIT = "CATEGORY_EDIT"
    CATEGORY_DELETE = "CATEGORY_DELETE"
    CATEGORY_SYNC = "CATEGORY_SYNC"

    LICENSE_CONTENT_ADD = "LICENSE_CONTENT_ADD"
    LICENSE_CONTENT_REMOVE = "LICENSE_CONTENT_REMOVE"

    CONTRACT_ATTACHMENT_DELETE = "CONTRACT_ATTACHMENT_DELETE"
    CONTRACT_ATTACHMENT_UPLOAD = "CONTRACT_ATTACHMENT_UPLOAD"

    GENRE_CREATE = "GENRE_CREATE"
    GENRE_EDIT = "GENRE_EDIT"
    GENRE_DELETE = "GENRE_DELETE"

    TAG_CREATE = "TAG_CREATE"
    TAG_EDIT = "TAG_EDIT"
    TAG_DELETE = "TAG_DELETE"

    CUSTOM_TAG_CREATE = "CUSTOM_TAG_CREATE"
    CUSTOM_TAG_EDIT = "CUSTOM_TAG_EDIT"
    CUSTOM_TAG_DELETE = "CUSTOM_TAG_DELETE"

    CONTENT_TYPE_CREATE = "CONTENT_TYPE_CREATE"
    CONTENT_TYPE_EDIT = "CONTENT_TYPE_EDIT"
    CONTENT_TYPE_DELETE = "CONTENT_TYPE_DELETE"

    POSTER_SIZE_CREATE = "POSTER_SIZE_CREATE"
    POSTER_SIZE_EDIT = "POSTER_SIZE_EDIT"
    POSTER_SIZE_DELETE = "POSTER_SIZE_DELETE"

    CUSTOM_FIELD_CREATE = "CUSTOM_FIELD_CREATE"
    CUSTOM_FIELD_EDIT = "CUSTOM_FIELD_EDIT"
    CUSTOM_FIELD_DELETE = "CUSTOM_FIELD_DELETE"

    SENSITIVE_WORD_CREATE = "SENSITIVE_WORD_CREATE"
    SENSITIVE_WORD_EDIT = "SENSITIVE_WORD_EDIT"
    SENSITIVE_WORD_DELETE = "SENSITIVE_WORD_DELETE"
    SENSITIVE_WORD_STATUS = "SENSITIVE_WORD_STATUS"
    SENSITIVE_WORD_IMPORT = "SENSITIVE_WORD_IMPORT"
    SENSITIVE_WORD_EXPORT = "SENSITIVE_WORD_EXPORT"

    TASK_ASSIGN = "TASK_ASSIGN"
    TASK_COMPLETE = "TASK_COMPLETE"

    DATA_AUTH_AUTHORIZE = "DATA_AUTH_AUTHORIZE"
    DATA_AUTH_CLEAR = "DATA_AUTH_CLEAR"

    METADATA_SOURCE_CREATE = "METADATA_SOURCE_CREATE"
    METADATA_SOURCE_EDIT = "METADATA_SOURCE_EDIT"
    METADATA_SOURCE_DELETE = "METADATA_SOURCE_DELETE"
    METADATA_SOURCE_STATUS = "METADATA_SOURCE_STATUS"

    CRAWL_TASK_CREATE = "CRAWL_TASK_CREATE"
    CRAWL_TASK_RETRY = "CRAWL_TASK_RETRY"
    CRAWL_TASK_DELETE = "CRAWL_TASK_DELETE"
    CRAWL_TASK_CONFIRM = "CRAWL_TASK_CONFIRM"

    MENU_CREATE = "MENU_CREATE"
    MENU_EDIT = "MENU_EDIT"
    MENU_DELETE = "MENU_DELETE"
    MENU_ASSIGN = "MENU_ASSIGN"

    WORKFLOW_CREATE = "WORKFLOW_CREATE"
    WORKFLOW_EDIT = "WORKFLOW_EDIT"
    WORKFLOW_DELETE = "WORKFLOW_DELETE"
    WORKFLOW_PUBLISH = "WORKFLOW_PUBLISH"
    WORKFLOW_UNPUBLISH = "WORKFLOW_UNPUBLISH"
    WORKFLOW_NEW_VERSION = "WORKFLOW_NEW_VERSION"
    WORKFLOW_BATCH_PUBLISH = "WORKFLOW_BATCH_PUBLISH"

    METADATA_QUALITY_TRIGGER = "METADATA_QUALITY_TRIGGER"
    METADATA_QUALITY_DELETE = "METADATA_QUALITY_DELETE"
    METADATA_QUALITY_EXPORT = "METADATA_QUALITY_EXPORT"

    VALIDATION_RULE_CREATE = "VALIDATION_RULE_CREATE"
    VALIDATION_RULE_EDIT = "VALIDATION_RULE_EDIT"
    VALIDATION_RULE_DELETE = "VALIDATION_RULE_DELETE"
    VALIDATION_RULE_BATCH_DELETE = "VALIDATION_RULE_BATCH_DELETE"
    VALIDATION_RULE_IMPORT = "VALIDATION_RULE_IMPORT"

    SCHEDULED_TASK_TRIGGER = "SCHEDULED_TASK_TRIGGER"

    USAGE_LIMIT_UPDATE = "USAGE_LIMIT_UPDATE"

    DASHBOARD_CONFIG_UPDATE = "DASHBOARD_CONFIG_UPDATE"
    DASHBOARD_CONFIG_RESET = "DASHBOARD_CONFIG_RESET"

    OPERATION_LOG_EXPORT = "OPERATION_LOG_EXPORT"
    OPERATION_LOG_CLEAR = "OPERATION_LOG_CLEAR"

    PROVIDER_BATCH_DELETE = "PROVIDER_BATCH_DELETE"
    CONTRACT_BATCH_DELETE = "CONTRACT_BATCH_DELETE"
    LICENSE_BATCH_DELETE = "LICENSE_BATCH_DELETE"
    USER_BATCH_DELETE = "USER_BATCH_DELETE"
    USER_BATCH_STATUS = "USER_BATCH_STATUS"
    ROLE_BATCH_DELETE = "ROLE_BATCH_DELETE"
    ROLE_BATCH_STATUS = "ROLE_BATCH_STATUS"
    SENSITIVE_WORD_BATCH_DELETE = "SENSITIVE_WORD_BATCH_DELETE"
    SENSITIVE_WORD_BATCH_STATUS = "SENSITIVE_WORD_BATCH_STATUS"
    SENSITIVE_WORD_BATCH_EXPORT = "SENSITIVE_WORD_BATCH_EXPORT"
    METADATA_QUALITY_BATCH_DELETE = "METADATA_QUALITY_BATCH_DELETE"
    SCHEDULED_TASK_BATCH_TRIGGER = "SCHEDULED_TASK_BATCH_TRIGGER"
    SCHEDULED_TASK_UPDATE = "SCHEDULED_TASK_UPDATE"
    CRAWL_TASK_BATCH_DELETE = "CRAWL_TASK_BATCH_DELETE"
    CUSTOM_FIELD_BATCH_DELETE = "CUSTOM_FIELD_BATCH_DELETE"
    TAG_BATCH_DELETE = "TAG_BATCH_DELETE"
    CATEGORY_BATCH_DELETE = "CATEGORY_BATCH_DELETE"
    CUSTOM_TAG_BATCH_DELETE = "CUSTOM_TAG_BATCH_DELETE"
    POSTER_SIZE_BATCH_DELETE = "POSTER_SIZE_BATCH_DELETE"
    CAST_BATCH_DELETE = "CAST_BATCH_DELETE"
    CONTENT_TYPE_BATCH_DELETE = "CONTENT_TYPE_BATCH_DELETE"
    METADATA_SOURCE_BATCH_ENABLE = "METADATA_SOURCE_BATCH_ENABLE"
    METADATA_SOURCE_BATCH_DISABLE = "METADATA_SOURCE_BATCH_DISABLE"
    METADATA_SOURCE_BATCH_DELETE = "METADATA_SOURCE_BATCH_DELETE"
    GENRE_BATCH_DELETE = "GENRE_BATCH_DELETE"
    CONTENT_BATCH_DELETE = "CONTENT_BATCH_DELETE"
    CONTENT_BATCH_EXPORT = "CONTENT_BATCH_EXPORT"
    CONTENT_BATCH_IMPORT = "CONTENT_BATCH_IMPORT"
    PACKAGE_BATCH_DELETE = "PACKAGE_BATCH_DELETE"
    TASK_BATCH_ASSIGN = "TASK_BATCH_ASSIGN"
    SCHEDULE_BATCH_EXPORT = "SCHEDULE_BATCH_EXPORT"

# ── 写入日志 ────────────────────────────────────────────────

async def write_log(
    db: AsyncSession,
    *,
    user_id: int | None,
    user_name: str | None,
    operation_type: str,
    operation_object: str | None = None,
    operation_object_code: str | None = None,
    operation_object_params: dict | None = None,
    operation_content: str | None = None,
    operation_content_code: str | None = None,
    operation_content_params: dict | None = None,
    content_id: int | None = None,
    entity_type: str | None = None,
    entity_id: int | None = None,
    previous_value: str | None = None,
    updated_value: str | None = None,
    updated_value_json: str | None = None,
    ip_address: str | None = None,
    result: str = "success",
    error_message: str | None = None,
    error_message_code: str | None = None,
    error_message_params: dict | None = None,
) -> None:
    """写入操作日志，内部捕获所有异常，不影响主业务流程。

    i18n 参数（可选）：
        operation_object_code / operation_object_params：
            若提供，用 encode_log_i18n 编码后作为 operation_object 入库。
        operation_content_code / operation_content_params：
            若提供，用 encode_log_i18n 编码后作为 operation_content 入库。
        error_message_code / error_message_params：
            若提供，用 encode_log_i18n 编码后作为 error_message 入库。
        若 *_code 为 None，则保持原有对应字段字符串（向后兼容旧调用）。

    值字段三值语义（统一约定）：
        previous_value / updated_value：富化后的展示值（外键已解析为名称、枚举已翻译），
            供详情页 Processed History 的 Previous / Updated value 两列展示。
        updated_value_json：原始数据快照（未富化、未翻译）——
            CREATE / DELETE 为实体全量快照（prepare_log_values 第三返回值）；
            UPDATE 为原始 diff（仅变更字段）；
            动作型 / 批量操作（无实体快照可查）由下方兜底逻辑自动落
            operation_content_params 序列化结果。
        注意：UPDATE 场景 updated_value_json 存的是 diff 而非全量，属约定行为，
        避免每次编辑都存全量快照导致存储膨胀。
    entity_type / entity_id 用于详情页历史精确查询。
    """
    try:
        if operation_object_code:
            operation_object = encode_log_i18n(operation_object_code, **(operation_object_params or {}))
        if operation_content_code:
            operation_content = encode_log_i18n(operation_content_code, **(operation_content_params or {}))
        if error_message_code:
            error_message = encode_log_i18n(error_message_code, **(error_message_params or {}))
        # 兜底：动作型/批量操作调用点通常不传值字段，原始入参已在
        # operation_content_params（names/count/ids 等），序列化后补进
        # updated_value_json，保证日志详情有据可查（超长截断保护）。
        if updated_value_json is None and updated_value is None and operation_content_params:
            updated_value_json = _json.dumps(
                operation_content_params, ensure_ascii=False, default=str
            )[:4000]
        if operation_object and len(operation_object) > 200:
            operation_object = operation_object[:197] + "..."
        log = OperationLog(
            user_id=user_id,
            user_name=user_name,
            operation_type=operation_type,
            operation_object=operation_object,
            operation_content=operation_content,
            previous_value=previous_value,
            updated_value=updated_value,
            updated_value_json=updated_value_json,
            content_id=content_id,
            entity_type=entity_type,
            entity_id=entity_id,
            ip_address=ip_address,
            result=result,
            error_message=error_message,
        )
        # 用 SAVEPOINT 隔离单条日志写入失败：仅回滚到保存点，保留事务中其余变更
        # （批量端点循环写 N 条日志时某条失败，不得连带回滚已写入的其他日志；
        #   亦避免 db.rollback() 回滚整个事务、吞掉尚未提交的业务变更）
        # 关键：db.add(log) 必须放在 begin_nested() 之内。若在 SAVEPOINT 之外 add，
        # flush 失败时异常会被 _capture_exception 传播到外层事务（_rollback_exception），
        # 使外层事务进入 DEACTIVE，后续 db.commit() 抛 PendingRollbackError——SAVEPOINT 隔离失效。
        # 放入 SAVEPOINT 内，回滚时会连同 pending 的 log 一并 expunge，外层事务保持可用。
        async with db.begin_nested():
            db.add(log)
            await db.flush()
    except Exception as exc:
        logger.warning("操作日志写入失败: %s", exc)

# ── 历史时间本地化（与导出 astimezone(app_tz) 保持一致） ───────
# 带时区偏移（Z 或 ±HH:MM）的 ISO 时间字符串/aware datetime 视为需要转换；
# naive 值（无偏移）按项目约定视为已是 app 本地时间，原样返回，避免双重偏移。
_ISO_DT_AWARE_RE = re.compile(
    r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}:\d{2}(?:\.\d+)?(?:Z|[+-]\d{2}:\d{2})$"
)


def _localize_dt(dt: datetime | None) -> datetime | None:
    """aware datetime → app 时区 naive 本地时间；naive 原样返回。"""
    if dt is None:
        return None
    if dt.tzinfo is not None:
        return dt.astimezone(app_tz).replace(tzinfo=None)
    return dt


def _localize_dt_str(value: str) -> str:
    """带时区偏移的 ISO 时间字符串 → app 时区 naive 本地字符串（'YYYY-MM-DD HH:MM:SS'）。

    无偏移的 naive 字符串（如纯日期、naive datetime）原样返回。
    """
    if not isinstance(value, str) or not _ISO_DT_AWARE_RE.match(value):
        return value
    try:
        dt = datetime.fromisoformat(value.replace(" ", "T").replace("Z", "+00:00"))
    except ValueError:
        return value
    if dt.tzinfo is None:
        return value
    return dt.astimezone(app_tz).replace(tzinfo=None).strftime("%Y-%m-%d %H:%M:%S")


def _localize_log_json(s: str | None) -> str | None:
    """递归把日志 JSON 字段里带偏移的时间字符串转为 app 时区本地字符串。

    非 JSON 字符串原样返回（updated_value 可能是文本摘要）。
    """
    if not s:
        return s
    try:
        obj = _json.loads(s)
    except (ValueError, TypeError):
        return s

    def _walk(v):
        if isinstance(v, str):
            return _localize_dt_str(v)
        if isinstance(v, dict):
            return {k: _walk(val) for k, val in v.items()}
        if isinstance(v, list):
            return [_walk(val) for val in v]
        return v

    return _json.dumps(_walk(obj), ensure_ascii=False)


# ── 通用实体历史查询 ──────────────────────────────────────────

async def list_entity_history(
    db: AsyncSession,
    entity_type: str,
    entity_id: int,
    limit: int = 100,
) -> list:
    from app.internal.cms_biz_system.schemas.operation_log import ProcessedHistoryItem
    logs = (
        await db.execute(
            select(OperationLog)
            .where(
                OperationLog.is_deleted.is_(False),
                OperationLog.entity_type == entity_type,
                OperationLog.entity_id == entity_id,
            )
            .order_by(OperationLog.operation_time.desc())
            .limit(limit)
        )
    ).scalars().all()

    user_ids = {log.user_id for log in logs if log.user_id}
    display_name_map: dict[int, str] = {}
    if user_ids:
        u_result = await db.execute(
            select(User.id, User.display_name).where(User.id.in_(user_ids))
        )
        display_name_map = {uid: dn for uid, dn in u_result.all() if dn}

    return [
        ProcessedHistoryItem(
            id=log.id,
            processed_at=_localize_dt(log.operation_time),
            processed_by=log.user_name,
            processed_by_display_name=display_name_map.get(log.user_id) if log.user_id else None,
            processed_type=log.operation_type,
            entity_type=log.entity_type,
            details=translate_log_value(log.operation_content),
            previous_value=_localize_log_json(log.previous_value),
            updated_value=_localize_log_json(log.updated_value),
            updated_value_json=_localize_log_json(log.updated_value_json),
        )
        for log in logs
    ]


async def list_content_history(
    db: AsyncSession,
    content_id: int,
    limit: int = 100,
) -> list:
    from app.internal.cms_biz_system.schemas.operation_log import ProcessedHistoryItem
    # Activity Log 仅展示当前内容自身的操作日志，不再聚合任何子内容日志：
    # 单季/单集的日志（含注入/移除等结构类操作）只归属子内容自身详情页，
    # 避免总季/单季详情页混入子内容记录（业务记录 bug 32142/32143，结构记录本次一并移除）
    logs = (
        await db.execute(
            select(OperationLog)
            .where(
                OperationLog.is_deleted.is_(False),
                OperationLog.content_id == content_id,
            )
            .order_by(OperationLog.operation_time.desc())
            .limit(limit)
        )
    ).scalars().all()

    user_ids = {log.user_id for log in logs if log.user_id}
    display_name_map: dict[int, str] = {}
    if user_ids:
        u_result = await db.execute(
            select(User.id, User.display_name).where(User.id.in_(user_ids))
        )
        display_name_map = {uid: dn for uid, dn in u_result.all() if dn}

    return [
        ProcessedHistoryItem(
            id=log.id,
            processed_at=_localize_dt(log.operation_time),
            processed_by=log.user_name,
            processed_by_display_name=display_name_map.get(log.user_id) if log.user_id else None,
            processed_type=log.operation_type,
            entity_type=log.entity_type,
            details=translate_log_value(log.operation_content),
            previous_value=_localize_log_json(log.previous_value),
            updated_value=_localize_log_json(log.updated_value),
            updated_value_json=_localize_log_json(log.updated_value_json),
        )
        for log in logs
    ]

# ── 查询 ─────────────────────────────────────────────────────

async def list_logs(
    db: AsyncSession,
    page: int = 1,
    page_size: int = 10,
    user_name: str | None = None,
    operation_type: str | None = None,
    operation_object: str | None = None,
    time_start: datetime | None = None,
    time_end: datetime | None = None,
    result: str | None = None,
    sort_by: str | None = None,
    sort_order: str | None = None,
) -> PaginatedResponse[OperationLogItem]:
    query = select(OperationLog)

    if user_name:
        query = query.where(OperationLog.user_name.ilike(f"%{user_name}%"))
    if operation_type:
        types = [t.strip() for t in operation_type.split(",") if t.strip()]
        if len(types) == 1:
            query = query.where(OperationLog.operation_type == types[0])
        else:
            query = query.where(OperationLog.operation_type.in_(types))
    if operation_object:
        query = query.where(OperationLog.operation_object.ilike(f"%{operation_object}%"))
    if time_start:
        query = query.where(OperationLog.operation_time >= time_start)
    if time_end:
        query = query.where(OperationLog.operation_time <= time_end)
    if result:
        query = query.where(OperationLog.result == result)

    total = (await db.execute(select(func.count()).select_from(query.subquery()))).scalar_one()

    if sort_by and sort_order:
        sort_column = getattr(OperationLog, sort_by, None)
        if sort_column is not None:
            order_func = sort_column.asc() if sort_order == 'asc' else sort_column.desc()
            query = query.order_by(order_func)
        else:
            query = query.order_by(OperationLog.operation_time.desc())
    else:
        query = query.order_by(OperationLog.operation_time.desc())

    logs = (
        await db.execute(
            query.offset((page - 1) * page_size)
            .limit(page_size)
        )
    ).scalars().all()

    user_ids = {log.user_id for log in logs if log.user_id}
    display_name_map: dict[int, str] = {}
    if user_ids:
        u_result = await db.execute(
            select(User.id, User.display_name).where(User.id.in_(user_ids))
        )
        display_name_map = {uid: dn for uid, dn in u_result.all() if dn}

    items = []
    for log in logs:
        item = OperationLogItem.model_validate(log)
        item.operation_object = translate_log_value(log.operation_object)
        item.operation_content = translate_log_value(log.operation_content)
        item.error_message = translate_log_value(log.error_message)
        item.user_display_name = display_name_map.get(log.user_id) if log.user_id else None
        items.append(item)

    return PaginatedResponse(
        total=total,
        page=page,
        page_size=page_size,
        items=items,
    )

async def get_log(db: AsyncSession, log_id: int) -> OperationLog:
    from fastapi import HTTPException, status as http_status
    log = (
        await db.execute(select(OperationLog).where(OperationLog.id == log_id))
    ).scalar_one_or_none()
    if not log:
        raise NotFoundException(ErrorCode.NOT_FOUND, get_msg("OPERATION_LOG_NOT_FOUND"))
    return log

# ── 清空 ─────────────────────────────────────────────────────

async def clear_logs(db: AsyncSession, start: datetime, end: datetime) -> int:
    result = await db.execute(
        delete(OperationLog).where(
            OperationLog.operation_time >= start,
            OperationLog.operation_time <= end,
        )
    )
    await db.commit()
    return result.rowcount

# ── Excel 导出 ────────────────────────────────────────────────

async def export_logs_excel(
    db: AsyncSession,
    user_name: str | None = None,
    operation_type: str | None = None,
    operation_object: str | None = None,
    time_start: datetime | None = None,
    time_end: datetime | None = None,
    result: str | None = None,
    ids: list[int] | None = None,
    lang: str = "cn",
) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from app.common.core.i18n import get_message

    query = select(OperationLog)
    if user_name:
        query = query.where(OperationLog.user_name.ilike(f"%{user_name}%"))
    if operation_type:
        types = [t.strip() for t in operation_type.split(",") if t.strip()]
        if len(types) == 1:
            query = query.where(OperationLog.operation_type == types[0])
        else:
            query = query.where(OperationLog.operation_type.in_(types))
    if operation_object:
        query = query.where(OperationLog.operation_object.ilike(f"%{operation_object}%"))
    if time_start:
        query = query.where(OperationLog.operation_time >= time_start)
    if time_end:
        query = query.where(OperationLog.operation_time <= time_end)
    if result:
        query = query.where(OperationLog.result == result)
    if ids:
        query = query.where(OperationLog.id.in_(ids))

    logs = (
        await db.execute(query.order_by(OperationLog.operation_time.desc()))
    ).scalars().all()

    user_ids = {log.user_id for log in logs if log.user_id}
    display_name_map: dict[int, str] = {}
    if user_ids:
        u_result = await db.execute(
            select(User.id, User.display_name).where(User.id.in_(user_ids))
        )
        display_name_map = {uid: dn for uid, dn in u_result.all() if dn}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = get_message("EXPORT_SHEET_TITLE", lang=lang)

    headers = [
        get_message("EXPORT_COL_USER", lang=lang),
        get_message("EXPORT_COL_TYPE", lang=lang),
        get_message("EXPORT_COL_OBJECT", lang=lang),
        get_message("EXPORT_COL_CONTENT", lang=lang),
        get_message("EXPORT_COL_TIME", lang=lang),
        get_message("EXPORT_COL_IP", lang=lang),
        get_message("EXPORT_COL_RESULT", lang=lang),
        get_message("EXPORT_COL_ERROR", lang=lang),
    ]
    header_fill = PatternFill(start_color="1677FF", end_color="1677FF", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    col_widths = [16, 18, 24, 40, 22, 16, 12, 30]
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    for row_idx, log in enumerate(logs, 2):
        op_time = log.operation_time.astimezone(app_tz).strftime("%Y-%m-%d %H:%M:%S") if log.operation_time else ""
        op_type = get_message(log.operation_type, lang=lang) if log.operation_type else ""
        dn = display_name_map.get(log.user_id) if log.user_id else None
        user_cell = f"{dn}({log.user_name})" if dn and log.user_name else (dn or log.user_name or "")
        ws.cell(row=row_idx, column=1, value=user_cell)
        ws.cell(row=row_idx, column=2, value=op_type)
        ws.cell(row=row_idx, column=3, value=_translate_for_export(log.operation_object, lang))
        # Content 列与前端列表保持一致：优先 updated_value_json，其次 updated_value，最后 operation_content（需翻译）
        content_value = log.updated_value_json or log.updated_value or _translate_for_export(log.operation_content, lang)
        ws.cell(row=row_idx, column=4, value=content_value)
        ws.cell(row=row_idx, column=5, value=op_time)
        ws.cell(row=row_idx, column=6, value=log.ip_address or "")
        ws.cell(row=row_idx, column=7, value=get_message("EXPORT_RESULT_SUCCESS" if log.result == "success" else "EXPORT_RESULT_FAILED", lang=lang))
        ws.cell(row=row_idx, column=8, value=translate_log_value(log.error_message) or "")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
