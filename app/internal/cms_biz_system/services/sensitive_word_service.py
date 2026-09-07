import io
from datetime import datetime

from fastapi import UploadFile, status
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from ..models.dict import DictNode
from ..models.sensitive_word import SensitiveWord
from app.common.core import SensitiveWordStatus
from app.common.core.exceptions import ErrorCode, BusinessException, NotFoundException
from app.internal.cms_biz_system.schemas.sensitive_word import (
    BatchStatusRequest,
    ImportResult,
    SensitiveWordCreate,
    SensitiveWordListItem,
    SensitiveWordUpdate,
)
from app.common.schemas import PaginatedResponse
from app.config import app_tz

from loguru import logger
from app.common.core.i18n import get_msg
from app.common.services.sensitive_check_service import SensitiveCheckService

async def list_sensitive_words(
    db: AsyncSession,
    page: int = 1,
    page_size: int = 10,
    keyword: str | None = None,
    type_codes: list[str] | None = None,
    status: str | None = None,
    created_start: datetime | None = None,
    created_end: datetime | None = None,
    sort_by: str | None = None,
    sort_order: str | None = None,
) -> PaginatedResponse[SensitiveWordListItem]:
    query = select(SensitiveWord).where(SensitiveWord.is_deleted == False)

    if keyword:
        query = query.where(SensitiveWord.keyword.ilike(f"%{keyword}%"))
    if type_codes:
        type_root = (await db.execute(
            select(DictNode.id).where(DictNode.parent_id.is_(None), DictNode.code == "Sensitive_Word_Type", DictNode.is_deleted == False)
        )).scalar_one_or_none()
        expanded_codes = set(type_codes)
        if type_root is not None:
            dict_names = (await db.execute(
                select(DictNode.code, DictNode.name).where(
                    DictNode.parent_id == type_root,
                    DictNode.code.in_(type_codes),
                    DictNode.status != "deleted",
                    DictNode.is_deleted == False,
                )
            )).all()
            for d_code, d_name in dict_names:
                expanded_codes.add(d_name)
        query = query.where(SensitiveWord.type_code.in_(expanded_codes))
    if status:
        query = query.where(SensitiveWord.status == status)
    if created_start:
        query = query.where(SensitiveWord.created_at >= created_start)
    if created_end:
        query = query.where(SensitiveWord.created_at <= created_end)

    if sort_by and sort_order:
        sort_column = getattr(SensitiveWord, sort_by, None)
        if sort_column is not None:
            order_func = sort_column.asc() if sort_order == "asc" else sort_column.desc()
            query = query.order_by(order_func, SensitiveWord.id.desc())
        else:
            query = query.order_by(SensitiveWord.id.desc())
    else:
        query = query.order_by(SensitiveWord.id.desc())

    count_result = await db.execute(select(func.count()).select_from(query.subquery()))
    total = count_result.scalar_one()

    items = (
        await db.execute(query.offset((page - 1) * page_size).limit(page_size))
    ).scalars().all()

    return PaginatedResponse(
        total=total,
        page=page,
        page_size=page_size,
        items=[SensitiveWordListItem.model_validate(item) for item in items],
    )

async def get_sensitive_word(db: AsyncSession, word_id: int) -> SensitiveWord:
    word = (
        await db.execute(
            select(SensitiveWord).where(
                SensitiveWord.id == word_id,
                SensitiveWord.is_deleted == False,
            )
        )
    ).scalar_one_or_none()
    if not word:
        raise NotFoundException(ErrorCode.SENSITIVE_WORD_NOT_FOUND, get_msg("SENSITIVE_WORD_NOT_FOUND"))
    return word

async def create_sensitive_word(
    db: AsyncSession, data: SensitiveWordCreate
) -> SensitiveWord:
    # 敏感词以"关键词"为唯一身份（与导入逻辑、唯一索引一致）。
    # 不过滤软删除行：命中软删除行时复活并更新（与导入路径行为一致），
    # 避免在仍保留旧版非部分唯一约束的环境（软删除行占用唯一键）上 INSERT 撞键导致 500。
    existing = (
        await db.execute(
            select(SensitiveWord)
            .where(SensitiveWord.keyword == data.keyword)
            .order_by(SensitiveWord.is_deleted.asc(), SensitiveWord.id.asc())
            .limit(1)
        )
    ).scalar_one_or_none()
    if existing and not existing.is_deleted:
        raise BusinessException(ErrorCode.SENSITIVE_WORD_EXISTS, get_msg("SENSITIVE_WORD_EXISTS"))
    if existing:
        # 复活软删除行，类型/状态以本次提交为准
        existing.type_code = data.type_code
        existing.status = data.status
        existing.is_deleted = False
        word = existing
    else:
        word = SensitiveWord(
            keyword=data.keyword,
            type_code=data.type_code,
            status=data.status,
        )
        db.add(word)
    await db.commit()
    await db.refresh(word)
    SensitiveCheckService.get_instance().invalidate_cache()
    return word

async def update_sensitive_word(
    db: AsyncSession, word_id: int, data: SensitiveWordUpdate
) -> SensitiveWord:
    word = await get_sensitive_word(db, word_id)
    new_keyword = data.keyword if data.keyword is not None else word.keyword
    new_type_code = data.type_code if data.type_code is not None else word.type_code

    if new_keyword != word.keyword:
        existing = (
            await db.execute(
                select(SensitiveWord).where(
                    SensitiveWord.keyword == new_keyword,
                    SensitiveWord.id != word_id,
                    SensitiveWord.is_deleted == False,
                )
            )
        ).scalar_one_or_none()
        if existing:
            raise BusinessException(ErrorCode.SENSITIVE_WORD_EXISTS, get_msg("SENSITIVE_WORD_EXISTS"))

    if data.keyword is not None:
        word.keyword = data.keyword
    if data.type_code is not None:
        word.type_code = data.type_code
    if data.status is not None:
        word.status = data.status
    await db.commit()
    await db.refresh(word)
    SensitiveCheckService.get_instance().invalidate_cache()
    return word

async def delete_sensitive_word(db: AsyncSession, word_id: int) -> None:
    word = await get_sensitive_word(db, word_id)
    word.is_deleted = True
    await db.commit()
    SensitiveCheckService.get_instance().invalidate_cache()

async def batch_delete(db: AsyncSession, ids: list[int]) -> int:
    words = (
        await db.execute(
            select(SensitiveWord).where(
                SensitiveWord.id.in_(ids),
                SensitiveWord.is_deleted == False,
            )
        )
    ).scalars().all()
    for word in words:
        word.is_deleted = True
    await db.commit()
    SensitiveCheckService.get_instance().invalidate_cache()
    return len(words)

async def toggle_status(db: AsyncSession, word_id: int, new_status: str) -> SensitiveWord:
    word = await get_sensitive_word(db, word_id)
    word.status = new_status
    await db.commit()
    await db.refresh(word)
    SensitiveCheckService.get_instance().invalidate_cache()
    return word

async def batch_toggle_status(db: AsyncSession, req: BatchStatusRequest) -> int:
    words = (
        await db.execute(
            select(SensitiveWord).where(
                SensitiveWord.id.in_(req.ids),
                SensitiveWord.is_deleted == False,
            )
        )
    ).scalars().all()
    for word in words:
        word.status = req.status
    await db.commit()
    SensitiveCheckService.get_instance().invalidate_cache()
    return len(words)

async def export_excel(db: AsyncSession, ids: list[int]) -> bytes:
    import openpyxl

    words = (
        await db.execute(
            select(SensitiveWord).where(
                SensitiveWord.id.in_(ids),
                SensitiveWord.is_deleted == False,
            )
        )
    ).scalars().all()

    # 查询 Sensitive_Word_Type 字典树，构建 type_code → type_name 映射
    type_root = (
        await db.execute(
            select(DictNode).where(
                DictNode.parent_id.is_(None),
                DictNode.code == "Sensitive_Word_Type",
                DictNode.is_deleted == False,
            )
        )
    ).scalar_one_or_none()
    type_name_map: dict[str, str] = {}
    if type_root is not None:
        children = (
            await db.execute(
                select(DictNode).where(
                    DictNode.parent_id == type_root.id,
                    DictNode.is_deleted == False,
                )
            )
        ).scalars().all()
        type_name_map = {c.code: c.name for c in children}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = get_msg("SENSITIVE_WORD_EXPORT_SHEET_TITLE")

    headers = [
        get_msg("SENSITIVE_WORD_EXPORT_COL_KEYWORD"),
        get_msg("SENSITIVE_WORD_EXPORT_COL_TYPE"),
        get_msg("SENSITIVE_WORD_EXPORT_COL_STATUS"),
        get_msg("SENSITIVE_WORD_EXPORT_COL_CREATED_AT"),
        get_msg("SENSITIVE_WORD_EXPORT_COL_UPDATED_AT"),
    ]
    header_fill = PatternFill(start_color="1677FF", end_color="1677FF", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    col_widths = [24, 20, 12, 22, 22]
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    for row_idx, word in enumerate(words, 2):
        created = word.created_at.astimezone(app_tz).strftime("%Y-%m-%d %H:%M:%S") if word.created_at else ""
        updated = word.updated_at.astimezone(app_tz).strftime("%Y-%m-%d %H:%M:%S") if word.updated_at else ""
        status_text = (
            get_msg("SENSITIVE_WORD_EXPORT_STATUS_ACTIVE")
            if word.status == "active"
            else get_msg("SENSITIVE_WORD_EXPORT_STATUS_INACTIVE")
        )
        ws.cell(row=row_idx, column=1, value=word.keyword)
        ws.cell(row=row_idx, column=2, value=type_name_map.get(word.type_code, word.type_code))
        ws.cell(row=row_idx, column=3, value=status_text)
        ws.cell(row=row_idx, column=4, value=created)
        ws.cell(row=row_idx, column=5, value=updated)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def generate_import_template() -> bytes:
    """生成导入模板（表头按请求语言本地化，列结构与 import_excel 解析一致）"""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = get_msg("SENSITIVE_WORD_EXPORT_SHEET_TITLE")

    headers = [
        get_msg("SENSITIVE_WORD_EXPORT_COL_KEYWORD"),
        get_msg("SENSITIVE_WORD_EXPORT_COL_TYPE"),
        get_msg("SENSITIVE_WORD_EXPORT_COL_STATUS"),
    ]
    header_fill = PatternFill(start_color="1677FF", end_color="1677FF", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for col, width in enumerate([24, 20, 12], 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()

async def _get_or_create_type_code(db: AsyncSession, type_name: str) -> str:
    root = (
        await db.execute(
            select(DictNode).where(
                DictNode.parent_id.is_(None),
                DictNode.code == "Sensitive_Word_Type",
                DictNode.is_deleted == False,
            )
        )
    ).scalar_one_or_none()
    if root is None:
        raise NotFoundException(ErrorCode.DICT_SENSITIVE_WORD_TYPE_NOT_FOUND, get_msg("DICT_SENSITIVE_WORD_TYPE_NOT_FOUND"))

    existing = (
        await db.execute(
            select(DictNode).where(
                DictNode.parent_id == root.id,
                DictNode.name == type_name,
                DictNode.status != "deleted",
                DictNode.is_deleted == False,
            )
        )
    ).scalar_one_or_none()
    if existing:
        return existing.code

    max_sort = (
        await db.execute(
            select(func.coalesce(func.max(DictNode.sort_order), -1)).where(
                DictNode.parent_id == root.id,
                DictNode.is_deleted == False,
            )
        )
    ).scalar_one()
    new_node = DictNode(
        parent_id=root.id,
        code=type_name,
        name=type_name,
        sort_order=max_sort + 1,
        status=SensitiveWordStatus.ACTIVE.value,
        is_system=False,
    )
    db.add(new_node)
    await db.flush()
    return new_node.code

async def import_excel(db: AsyncSession, file: UploadFile) -> ImportResult:
    content = await file.read()
    wb = load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active

    # 表头校验：兼容中英文模板（关键词/Keyword、类型/Type、状态/Status），按表头名定位列
    header_aliases = {
        "keyword": {"关键词", "Keyword"},
        "type": {"类型", "Type"},
        "status": {"状态", "Status"},
    }
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        wb.close()
        raise BusinessException(ErrorCode.INVALID_FILE_FORMAT, get_msg("INVALID_FILE_FORMAT"))
    actual_headers = [str(c).strip() if c is not None else "" for c in header_row]
    col_index: dict[str, int] = {}
    for field, aliases in header_aliases.items():
        for idx, h in enumerate(actual_headers):
            if h in aliases:
                col_index[field] = idx
                break
    missing = [f for f in header_aliases if f not in col_index]
    if missing:
        wb.close()
        raise BusinessException(ErrorCode.INVALID_FILE_FORMAT, get_msg("INVALID_FILE_FORMAT"))

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()

    def _cell(row: tuple, field: str) -> str:
        idx = col_index.get(field)
        if idx is None or idx >= len(row):
            return ""
        v = row[idx]
        return str(v).strip() if v is not None else ""

    result = ImportResult()
    for row in rows:
        if not row:
            continue
        keyword = _cell(row, "keyword")
        type_name = _cell(row, "type")
        status_str = _cell(row, "status") or "active"

        if not keyword or not type_name:
            continue

        result.total += 1
        type_code = await _get_or_create_type_code(db, type_name)

        if status_str.lower() in ("已禁用", "inactive"):
            status_val = "inactive"
        else:
            status_val = "active"

        # 需求 3.1.6.2：敏感词以"关键词"为唯一身份，已存在则更新（类型/状态以最新导入为准）。
        # 不过滤软删除行：命中软删除行时复活并更新，避免撞 keyword 部分唯一索引；
        # is_deleted 升序保证优先命中未删除行（false < true）。
        existing = (
            await db.execute(
                select(SensitiveWord)
                .where(SensitiveWord.keyword == keyword)
                .order_by(SensitiveWord.is_deleted.asc(), SensitiveWord.id.asc())
                .limit(1)
            )
        ).scalar_one_or_none()

        if existing:
            existing.type_code = type_code
            existing.status = status_val
            existing.is_deleted = False
            result.updated += 1
        else:
            db.add(
                SensitiveWord(
                    keyword=keyword,
                    type_code=type_code,
                    status=status_val,
                )
            )
            result.created += 1

    await db.commit()
    SensitiveCheckService.get_instance().invalidate_cache()
    return result
