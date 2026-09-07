"""
元数据质量检查（MetadataQuality）服务层。

职责：
- 列表分页查询（支持时间区间、status 多选）；
- 详情（含 issues）；
- 手动触发：生成 pending 记录 → 后台执行 MetadataQualityCheck 空壳业务逻辑；
- 批量软删除；
- Excel 报告导出。
"""

import io
from datetime import datetime

from app.common.core.exceptions import BusinessException, ErrorCode
from app.common.core.i18n import get_msg
# 异常处理已内联到具体函数中
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import asc, desc, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.common.schemas import PaginatedResponse
from app.config import app_tz
from app.internal.cms_biz_system.models.metadata_quality import (
    MetadataQualityCheck,
    MetadataQualityIssue,
)
from app.internal.cms_biz_system.models.scheduled_task import ScheduledTask
from app.internal.cms_biz_system.schemas.metadata_quality import (
    MetadataQualityCheckDetail,
    MetadataQualityCheckOut,
    MetadataQualityIssueOut,
)

_SORT_FIELDS = {"id", "check_time", "status", "total_contents", "passed_count", "failed_count", "duration"}


async def list_metadata_quality_checks(
    db: AsyncSession,
    page: int = 1,
    page_size: int = 10,
    time_start: datetime | None = None,
    time_end: datetime | None = None,
    status_list: list[str] | None = None,
    sort_by: str | None = None,
    sort_order: str | None = None,
) -> PaginatedResponse[MetadataQualityCheckOut]:
    query = select(MetadataQualityCheck).where(MetadataQualityCheck.is_deleted.is_(False))

    if status_list:
        query = query.where(MetadataQualityCheck.status.in_(status_list))
    if time_start:
        query = query.where(MetadataQualityCheck.check_time >= time_start)
    if time_end:
        query = query.where(MetadataQualityCheck.check_time <= time_end)

    if sort_by and sort_by in _SORT_FIELDS:
        col = getattr(MetadataQualityCheck, sort_by)
        query = query.order_by(asc(col) if sort_order == "asc" else desc(col), MetadataQualityCheck.id.desc())
    else:
        query = query.order_by(MetadataQualityCheck.check_time.desc(), MetadataQualityCheck.id.desc())

    total = (await db.execute(select(func.count()).select_from(query.subquery()))).scalar_one()
    rows = (
        await db.execute(query.offset((page - 1) * page_size).limit(page_size))
    ).scalars().all()

    return PaginatedResponse(
        total=total,
        page=page,
        page_size=page_size,
        items=[MetadataQualityCheckOut.model_validate(r) for r in rows],
    )


_ISSUE_SORT_FIELDS = {"id", "content_id", "content_name", "content_type", "issue_type", "field_name", "severity"}


async def list_metadata_quality_issues(
    db: AsyncSession,
    check_id: int,
    page: int = 1,
    page_size: int = 10,
    sort_by: str | None = None,
    sort_order: str | None = None,
) -> PaginatedResponse[MetadataQualityIssueOut]:
    query = select(MetadataQualityIssue).where(MetadataQualityIssue.check_id == check_id)

    if sort_by and sort_by in _ISSUE_SORT_FIELDS:
        col = getattr(MetadataQualityIssue, sort_by)
        query = query.order_by(asc(col) if sort_order == "asc" else desc(col), MetadataQualityIssue.id.asc())
    else:
        query = query.order_by(MetadataQualityIssue.content_id.asc(), MetadataQualityIssue.id.asc())

    total = (await db.execute(select(func.count()).select_from(query.subquery()))).scalar_one()
    rows = (
        await db.execute(query.offset((page - 1) * page_size).limit(page_size))
    ).scalars().all()
    return PaginatedResponse(
        total=total,
        page=page,
        page_size=page_size,
        items=[MetadataQualityIssueOut.model_validate(r) for r in rows],
    )


async def get_metadata_quality_check(db: AsyncSession, check_id: int) -> MetadataQualityCheckDetail:
    check = (
        await db.execute(
            select(MetadataQualityCheck).where(
                MetadataQualityCheck.id == check_id,
                MetadataQualityCheck.is_deleted.is_(False),
            )
        )
    ).scalar_one_or_none()
    if check is None:
        raise BusinessException(ErrorCode.METADATA_QUALITY_CHECK_NOT_FOUND, get_msg("METADATA_QUALITY_CHECK_NOT_FOUND"))

    issues = (
        await db.execute(
            select(MetadataQualityIssue)
            .where(MetadataQualityIssue.check_id == check_id)
            .order_by(MetadataQualityIssue.id.asc())
        )
    ).scalars().all()

    base = MetadataQualityCheckOut.model_validate(check).model_dump()
    base["issues"] = [MetadataQualityIssueOut.model_validate(i) for i in issues]
    return MetadataQualityCheckDetail.model_validate(base)


async def trigger_metadata_quality_check(
    db: AsyncSession,
    operator: str | None = None,
) -> int:
    """
    手动触发一次元数据质检。

    - 前置校验：任务不存在/未启用/正在执行时直接抛业务异常拒绝，
      避免产生永远不会被推进的 pending 记录；
    - 校验通过后写入一条 status=pending 记录（便于前端列表立刻出现新项）；
    - 后台异步执行质检并把该记录推进到 completed/failed。
    """
    # 前置校验：任务必须存在且启用，且当前不在执行中（复用定时任务的通用错误码）
    task = (
        await db.execute(
            select(ScheduledTask).where(ScheduledTask.task_type == "MetadataQualityCheck")
        )
    ).scalar_one_or_none()
    if task is None or task.schedule_status != "enabled":
        raise BusinessException(
            ErrorCode.SCHEDULED_TASK_DISABLED_BLOCKED,
            get_msg("SCHEDULED_TASK_DISABLED_BLOCKED", task_types=["MetadataQualityCheck"]),
        )
    if task.execution_status == "running":
        raise BusinessException(
            ErrorCode.SCHEDULED_TASK_RUNNING_BLOCKED,
            get_msg("SCHEDULED_TASK_RUNNING_BLOCKED", task_types=["MetadataQualityCheck"]),
        )

    # 注意：模型里 created_by 是 int 外键（cms_user.id），而传入的 operator 是 username 字符串，
    # 这里不回填 created_by，保持为 NULL；operator 仅用于 trigger_task_manual 的操作人标记。
    record = MetadataQualityCheck(
        status="pending",
        total_contents=0,
        passed_count=0,
        failed_count=0,
        duration=None,
        is_deleted=False,
    )
    db.add(record)
    await db.commit()
    await db.refresh(record)

    # 触发后台任务，传入 check_id 以便更新 pending 记录
    from app.jobs.scheduler import trigger_task_manual

    await trigger_task_manual(db, "MetadataQualityCheck", operator=operator, check_id=record.id)

    return record.id


async def delete_metadata_quality_checks(db: AsyncSession, ids: list[int]) -> int:
    if not ids:
        return 0
    rows = (
        await db.execute(
            select(MetadataQualityCheck).where(
                MetadataQualityCheck.id.in_(ids),
                MetadataQualityCheck.is_deleted.is_(False),
            )
        )
    ).scalars().all()
    for row in rows:
        row.is_deleted = True
    await db.commit()
    return len(rows)


async def export_metadata_quality_report(db: AsyncSession, check_id: int) -> tuple[bytes, str]:
    """
    导出指定检查记录的 Excel 报告。
    Returns:
        (bytes, filename)
    """
    detail = await get_metadata_quality_check(db, check_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    header_fill = PatternFill(start_color="1677FF", end_color="1677FF", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    center = Alignment(horizontal="center", vertical="center")

    # Summary sheet
    summary_rows = [
        ("Check ID", detail.id),
        ("Check Time", detail.check_time.astimezone(app_tz).strftime("%Y-%m-%d %H:%M:%S")),
        ("Status", detail.status),
        ("Total Contents", detail.total_contents),
        ("Passed", detail.passed_count),
        ("Failed", detail.failed_count),
        ("Duration(s)", float(detail.duration) if detail.duration is not None else "-"),
    ]
    for idx, (k, v) in enumerate(summary_rows, 1):
        ws.cell(row=idx, column=1, value=k).font = Font(bold=True)
        ws.cell(row=idx, column=2, value=v)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 40

    # Issues sheet
    ws2 = wb.create_sheet("Issues")
    headers = ["Content ID", "Content Name", "Content Type", "Issue Type",
               "Field Name", "Severity", "Expected", "Actual"]
    for col, h in enumerate(headers, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
    for col_idx, width in enumerate([12, 32, 14, 12, 18, 10, 24, 24], 1):
        ws2.column_dimensions[ws2.cell(row=1, column=col_idx).column_letter].width = width

    for row_idx, issue in enumerate(detail.issues, 2):
        ws2.cell(row=row_idx, column=1, value=issue.content_id)
        ws2.cell(row=row_idx, column=2, value=issue.content_name)
        ws2.cell(row=row_idx, column=3, value=issue.content_type)
        ws2.cell(row=row_idx, column=4, value=issue.issue_type)
        ws2.cell(row=row_idx, column=5, value=issue.field_name)
        ws2.cell(row=row_idx, column=6, value=issue.severity)
        ws2.cell(row=row_idx, column=7, value=issue.expected_value)
        ws2.cell(row=row_idx, column=8, value=issue.actual_value)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read(), f"metadata_quality_report_{check_id}.xlsx"
