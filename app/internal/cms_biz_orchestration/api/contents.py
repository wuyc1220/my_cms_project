"""
内容（Content）API 路由层 — 交易管理视角。

路由前缀：/contents

接口列表：
    GET    /                         查询内容列表（分页 + 多维过滤）
    POST   /                         新建内容（含 SERIES/SEASON 子节点自动创建）
    GET    /without-license-count    统计无许可证内容数量（快捷按钮数据）
    GET    /series-simple            获取 SERIES 简要列表（EPISODE 父级下拉）
    GET    /channels-simple          获取 CHANNEL 简要列表（SCHEDULE 频道下拉）
    GET    /template/{content_type}  下载导入模板（EPISODE/SERIES）
    POST   /parse-excel              解析 Excel 导入文件（EPISODE/SERIES）
    GET    /{content_id}             查询单个内容详情
    PUT    /{content_id}             编辑内容
    DELETE /{content_id}             软删除内容（级联删除子节点）
    GET    /{content_id}/licenses    查询内容关联许可证列表
    GET    /{content_id}/adjacent    查询上一条/下一条内容 ID
"""

import io

from fastapi import APIRouter, Depends, Query, Request, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.common.dependencies import get_current_user, _get_ip
from app.common.dependencies import get_db
from app.common.utils.log_enricher import orm_to_dict, prepare_log_values, resolve_option_display_name
from app.internal.cms_biz_system.models.user import User
from app.internal.cms_biz_package.models.package import Content
from app.internal.cms_biz_orchestration.schemas.content import (
    AdjacentContentResponse,
    BatchImportRequest,
    BatchImportResponse,
    ContentCreate,
    ContentListItem,
    ContentSimpleItem,
    ContentUpdate,
    ContentLicenseRef,
    ContentDetailResponse,
)
from app.common.schemas import PaginatedResponse, BatchDeleteRequest
from app.internal.cms_biz_orchestration.services import content_service
from app.internal.cms_biz_orchestration.services import episode_history_service
from app.internal.cms_biz_orchestration.services import batch_content_service
from app.internal.cms_biz_orchestration.services import live_service
from app.internal.cms_biz_system.services.operation_log_service import OperationType, write_log
from app.internal.cms_biz_package.models.enums import ContentType
from app.internal.cms_biz_package.models.package import ContentGenre, ContentCustomTag
from app.internal.cms_biz_metada.models.basic import Genre, CustomTag
from app.internal.cms_biz_metada.services.entity_data_service import (
    get_field_values,
    save_field_values,
    get_i18n_values,
    save_i18n_values,
)
from app.internal.cms_biz_metada.schemas.basic import (
    EntityFieldValueItem,
    EntityFieldValuesPayload,
    EntityI18nItem,
    EntityI18nPayload,
)

router = APIRouter(prefix="/contents")


async def _load_genre_tag_names(db: AsyncSession, content_id: int) -> tuple[list[str], list[str]]:
    """查询内容关联的题材/自定义标签名称列表。

    题材与标签存于中间表（content_genre / content_custom_tag），不在 Content 主表列中，
    orm_to_dict 快照无法覆盖，编辑日志需手动补记名称供 diff 感知（对齐 metadata.py 的
    custom_tag_names 补记模式，bug 32430）。
    """
    genre_names = (await db.execute(
        select(Genre.name)
        .join(ContentGenre, ContentGenre.genre_id == Genre.id)
        .where(ContentGenre.content_id == content_id, Genre.is_deleted.is_(False))
    )).scalars().all()
    tag_names = (await db.execute(
        select(CustomTag.name)
        .join(ContentCustomTag, ContentCustomTag.custom_tag_id == CustomTag.id)
        .where(ContentCustomTag.content_id == content_id, CustomTag.is_deleted.is_(False))
    )).scalars().all()
    return list(genre_names), list(tag_names)


@router.get("/", response_model=PaginatedResponse[ContentListItem])
async def get_content_list(
    page: int = 1,
    page_size: int = 10,
    content_id: int | None = None,
    external_id: str | None = None,
    title: str | None = None,
    content_types: list[str] | None = Query(default=None),
    statuses: list[str] | None = Query(default=None),
    genre_ids: list[int] | None = Query(default=None),
    custom_tag_ids: list[int] | None = Query(default=None),
    parent_id: int | None = None,
    created_from: str | None = None,
    created_to: str | None = None,
    without_license: bool = False,
    license_start_from: str | None = None,
    license_start_to: str | None = None,
    license_end_from: str | None = None,
    license_end_to: str | None = None,
    is_discarded: bool | None = None,
    sort_by: str | None = None,
    sort_order: str | None = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return await content_service.list_contents(
        db, page, page_size, content_id, external_id, title,
        content_types, statuses, genre_ids, custom_tag_ids, parent_id,
        created_from, created_to,
        without_license,
        license_start_from, license_start_to,
        license_end_from, license_end_to,
        is_discarded=is_discarded,
        sort_by=sort_by, sort_order=sort_order,
        current_user=current_user,
    )


@router.post("/", response_model=ContentListItem)
async def create_content_api(
    body: ContentCreate,
    request: Request,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    content = await content_service.create_content(db, body, processed_by=current_user.username)

    # 如果创建的是 EPISODE 或 SERIES 类型，记录操作历史
    if content.content_type in ("EPISODE", "SERIES", "SEASON_SERIES") and content.parent_id:
        processed_by = current_user.username
        series_ordinal = content.series_ordinal if content.content_type in ("SERIES", "SEASON_SERIES") else None

        await episode_history_service.add_episode_history(
            db,
            parent_id=content.parent_id,
            content_id=content.id,
            content_name=content.title,
            content_type=content.content_type,
            series_ordinal=series_ordinal,
            processed_by=processed_by,
            processed_type="Add",
            created_by=current_user.id,
        )

    # 新增子内容后回退祖先状态（与删除对称：父级构成变化需重新审核）
    if content.parent_id:
        from app.internal.cms_biz_orchestration.services.workflow_service import rollback_ancestors_after_child_change
        child_type_label = {"EPISODE": "单集", "SERIES": "单季", "SEASON_SERIES": "单季", "SCHEDULE": "节目单"}.get(
            content.content_type, content.content_type
        )
        await rollback_ancestors_after_child_change(
            db,
            start_parent_id=content.parent_id,
            edited_by=current_user.username,
            edit_info=f"新增{child_type_label}「{content.title}」",
        )

    if content.content_type == ContentType.SCHEDULE.value:
        # 新增节目单日志与 /live/schedules（create_schedule_api）完全同构：
        # SCHEDULE_CREATE + entity_type="schedule" + 节目单口径快照（关联频道语义）。
        # 不沿用 ContentListItem 快照，避免混入 parent_id/parent_title（父子语义）
        # 与 license_count/task_start_time 等列表派生噪音字段
        schedule_orm = (
            await db.execute(select(Content).where(Content.id == content.id))
        ).scalar_one()
        schedule_item = await live_service.build_schedule_item(db, schedule_orm)
        new_data = orm_to_dict(schedule_item, "schedule")
        # 题材/自定义标签存于中间表，ScheduleListItem 快照不覆盖，
        # 新增日志需手动补记名称（与编辑路径 _load_genre_tag_names 补记模式一致，
        # 否则新增节目单填写题材/标签在 Activity Log 中不可见）
        genre_names, tag_names = await _load_genre_tag_names(db, content.id)
        new_data["genre_names"] = genre_names
        new_data["custom_tag_names"] = tag_names
        prev_val, upd_val, raw_val = await prepare_log_values(db, "schedule", None, new_data)
        await write_log(
            db,
            user_id=current_user.id,
            user_name=current_user.username,
            operation_type=OperationType.SCHEDULE_CREATE,
            operation_object_code="OBJ_SCHEDULE", operation_object_params={"name": content.title},
            operation_content_code="log.metadata.create",
            content_id=content.id,
            entity_type="schedule",
            entity_id=content.id,
            previous_value=prev_val,
            updated_value=upd_val,
            updated_value_json=raw_val,
            ip_address=_get_ip(request),
            result="success",
        )
    else:
        new_data = orm_to_dict(content, "content")
        # 题材/自定义标签存于中间表，Content 主表快照不覆盖，补记名称
        # （新建弹窗对所有内容类型均提供题材/标签输入，与上方 SCHEDULE 分支
        # 及编辑路径 _load_genre_tag_names 补记模式一致，否则 Add 日志不可见）
        genre_names, tag_names = await _load_genre_tag_names(db, content.id)
        new_data["genre_names"] = genre_names
        new_data["custom_tag_names"] = tag_names
        _, upd_val, raw_val = await prepare_log_values(db, "content", None, new_data)
        if content.content_type == ContentType.EPISODE.value:
            # 单集注入日志归属父级（content_id=parent_id），与批量导入（batch_create_contents_api）口径一致，
            # 保证父级（SERIES/SEASON_SERIES）详情页 Activity Log 可见；
            # entity_id 仍指向新 EPISODE 自身，保留实体定位能力
            await write_log(
                db,
                user_id=current_user.id,
                user_name=current_user.username,
                operation_type=OperationType.EPISODE_INJECT,
                operation_object_code="OBJ_CONTENT", operation_object_params={"name": content.title},
                operation_content_code="log.episode.inject", operation_content_params={"title": content.title or ""},
                content_id=content.parent_id,
                entity_type="content",
                entity_id=content.id,
                updated_value=upd_val,
                updated_value_json=raw_val,
                ip_address=_get_ip(request),
                result="success",
            )
        elif content.content_type == ContentType.SEASON_SERIES.value and content.parent_id:
            # 季详情页 Season Series 弹窗注入单季系列：日志归属父级（content_id=parent_id），
            # 与 EPISODE_INJECT 口径一致，保证父级（SEASON）详情页 Activity Log 可见；
            # entity_id 仍指向新 SEASON_SERIES 自身，保留实体定位能力
            await write_log(
                db,
                user_id=current_user.id,
                user_name=current_user.username,
                operation_type=OperationType.SEASON_SERIES_INJECT,
                operation_object_code="OBJ_CONTENT", operation_object_params={"name": content.title},
                operation_content_code="log.seasonSeries.inject", operation_content_params={"title": content.title or ""},
                content_id=content.parent_id,
                entity_type="content",
                entity_id=content.id,
                updated_value=upd_val,
                updated_value_json=raw_val,
                ip_address=_get_ip(request),
                result="success",
            )
        else:
            await write_log(
                db,
                user_id=current_user.id,
                user_name=current_user.username,
                operation_type=OperationType.CONTENT_CREATE,
                operation_object_code="OBJ_CONTENT", operation_object_params={"name": content.title},
                operation_content_code="LOG_CONTENT_CREATE", operation_content_params={"title": content.title},
                content_id=content.id,
                entity_type="content",
                entity_id=content.id,
                updated_value=upd_val,
                updated_value_json=raw_val,
                ip_address=_get_ip(request),
                result="success",
            )
    await db.commit()
    return content


@router.post("/batch", response_model=BatchImportResponse)
async def batch_create_contents_api(
    body: BatchImportRequest,
    request: Request,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """批量导入子内容（EPISODE/SERIES），单个事务控制。"""
    try:
        results, parent_ctype = await batch_content_service.batch_create_sub_contents(
            db, parent_id=body.parent_id, items=body.items, current_user=current_user,
        )

        success_count = sum(1 for r in results if r.success)
        failed_count = sum(1 for r in results if not r.success)

        import json
        raw_val = json.dumps({
            "parent_id": body.parent_id,
            "success_count": success_count,
            "failed_count": failed_count,
            "results": [{"title": r.title, "success": r.success, "message": r.error or ""} for r in results]
        }, ensure_ascii=False)

        await write_log(
            db,
            user_id=current_user.id,
            user_name=current_user.username,
            operation_type=OperationType.CONTENT_CREATE,
            operation_object_code="OBJ_BATCH_CREATE", operation_object_params={"name": f"{success_count + failed_count} 个子内容"},
            operation_content_code="LOG_CONTENT_CREATE", operation_content_params={"title": ", ".join(r.title for r in results if r.success and r.title)},
            content_id=body.parent_id,
            entity_type="content",
            entity_id=body.parent_id,
            previous_value=None,
            updated_value=f"批量导入 {success_count} 个内容",
            updated_value_json=raw_val,
            ip_address=_get_ip(request),
            result="success" if failed_count == 0 else "partial",
        )

        # 统一提交：内容 + 历史 + 日志 都在同一个事务中
        await db.commit()

        return BatchImportResponse(
            success_count=success_count,
            failed_count=failed_count,
            details=results,
        )
    except Exception:
        await db.rollback()
        raise


@router.get("/without-license-count")
async def get_without_license_count(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    count = await content_service.get_without_license_count(db, current_user=current_user)
    return {"count": count}


@router.get("/series-simple", response_model=list[ContentSimpleItem])
async def get_series_simple(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return await content_service.get_series_simple(db, current_user=current_user)


@router.get("/seasons-simple", response_model=list[ContentSimpleItem])
async def get_seasons_simple(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return await content_service.get_seasons_simple(db, current_user=current_user)


@router.get("/channels-simple", response_model=list[ContentSimpleItem])
async def get_channels_simple(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return await content_service.get_channels_simple(db, current_user=current_user)


@router.get("/{content_id}", response_model=ContentDetailResponse)
async def get_content_detail(
    content_id: int,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(get_current_user),
):
    return await content_service.get_content(db, content_id)


@router.put("/{content_id}", response_model=ContentListItem)
async def update_content_api(
    content_id: int,
    body: ContentUpdate,
    request: Request,
    skip_metadata_process: bool = False,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    old_detail = await content_service.get_content(db, content_id)
    old = old_detail.content
    # 获取 update 之前的 Content ORM 快照（用于日志比较，避免 ContentListItem 额外字段干扰）
    old_orm = (
        await db.execute(select(Content).where(Content.id == content_id))
    ).scalar_one()
    # 在 update 之前立即计算 old_data 快照，避免 SQLAlchemy identity map 导致 old_orm 被后续修改
    old_data = orm_to_dict(old_orm, "content")
    # 题材/标签存中间表，主表快照不覆盖，补记名称使 diff 可感知变更（否则仅改题材/标签时不产生任何日志）
    old_genre_names, old_tag_names = await _load_genre_tag_names(db, content_id)
    old_data["genre_names"] = old_genre_names
    old_data["custom_tag_names"] = old_tag_names
    if old.content_type == ContentType.SCHEDULE.value:
        # 节目单所属频道存于 parent_id，补记 channel_id 供富化为频道名称（对齐创建日志 schedule 口径）
        old_data["channel_id"] = old_orm.parent_id
    # 已发布/已下架内容编辑后回滚状态并新建提交审核记录（与 Metadata 编辑行为一致）
    from app.internal.cms_biz_orchestration.services.workflow_service import rollback_after_published_edit
    await rollback_after_published_edit(
        db,
        content_id=content_id,
        content_type=old.content_type,
        edited_by=current_user.username,
        edit_info="编辑内容基础信息",
    )
    content = await content_service.update_content(db, content_id, body)
    # 获取 update 之后的 Content ORM 快照
    new_orm = (
        await db.execute(select(Content).where(Content.id == content_id))
    ).scalar_one()
    new_data = orm_to_dict(new_orm, "content")
    new_genre_names, new_tag_names = await _load_genre_tag_names(db, content_id)
    new_data["genre_names"] = new_genre_names
    new_data["custom_tag_names"] = new_tag_names
    if old.content_type == ContentType.SCHEDULE.value:
        new_data["channel_id"] = new_orm.parent_id
    prev_val, upd_val, raw_val = await prepare_log_values(db, "content", old_data, new_data)
    # 仅当 Content 主表字段有实际变化时才记录日志，避免修改元数据/字段值时产生重复日志
    if prev_val is not None or upd_val is not None:
        # 频道/节目单基础信息编辑同样需要记录日志（此前提前 return 导致漏记）
        if old.content_type == ContentType.CHANNEL.value:
            op_type = OperationType.CHANNEL_UPDATE
            obj_code, content_code = "OBJ_CHANNEL", "log.channel.edit"
        elif old.content_type == ContentType.SCHEDULE.value:
            op_type = OperationType.SCHEDULE_UPDATE
            obj_code, content_code = "OBJ_SCHEDULE", "log.metadata.edit"
        else:
            op_type = OperationType.CONTENT_EDIT
            obj_code, content_code = "OBJ_CONTENT", "LOG_CONTENT_EDIT"
        await write_log(
            db,
            user_id=current_user.id,
            user_name=current_user.username,
            operation_type=op_type,
            operation_object_code=obj_code, operation_object_params={"name": old.title},
            operation_content_code=content_code, operation_content_params={"title": old.title},
            content_id=content_id,
            entity_type="content",
            entity_id=content_id,
            previous_value=prev_val,
            updated_value=upd_val,
            updated_value_json=raw_val,
            ip_address=_get_ip(request),
            result="success",
        )
        # 节目单/频道基础信息（分类/标签/起止时间）属于 Metadata 节点数据，存于 content 主表：
        # 独立编辑基础信息（EditContentModal）时补写流程记录保证 Processes 页签可见；
        # 元数据弹窗链路（MetadataModal 保存时同步题材/标签）传 skip_metadata_process=true 跳过，
        # 此时元数据尚未保存 → check_metadata_complete=False 会写入 Pending 中间记录，
        # 最终由 create/update_xxx_metadata 统一写入一条（bug：Processes 出现两条 Metadata）
        if old.content_type in (ContentType.SCHEDULE.value, ContentType.CHANNEL.value) and not skip_metadata_process:
            from app.internal.cms_biz_orchestration.services.workflow_service import complete_process_and_update_status
            await complete_process_and_update_status(
                db,
                content_id=content_id,
                content_type=old.content_type,
                process_name="Metadata",
                processed_by=current_user.username,
                info="修改基础信息（分类/标签/起止时间）",
            )
    await db.commit()
    return content


@router.delete("/{content_id}")
async def delete_content_api(
    content_id: int,
    request: Request,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    content_detail = await content_service.get_content(db, content_id)
    content_title = content_detail.content.title
    content_type = content_detail.content.content_type
    parent_id = content_detail.content.parent_id

    # 如果删除的是 EPISODE 或 SERIES/SEASON_SERIES 类型，记录操作历史（在删除前记录）
    if content_type in ("EPISODE", "SERIES", "SEASON_SERIES") and parent_id:
        processed_by = current_user.username

        # 如果是 SERIES/SEASON_SERIES 类型，需要获取 series_ordinal
        series_ordinal = None
        if content_type in ("SERIES", "SEASON_SERIES"):
            series_ordinal = getattr(content_detail.content, 'series_ordinal', None)

        await episode_history_service.add_episode_history(
            db,
            parent_id=parent_id,
            content_id=content_id,
            content_name=content_title,
            content_type=content_type,
            series_ordinal=series_ordinal,
            processed_by=processed_by,
            processed_type="Delete",
            created_by=current_user.id,
        )

    old_data = orm_to_dict(content_detail.content, "content")
    prev_val, _, raw_val = await prepare_log_values(db, "content", old_data, None)
    await content_service.delete_content(db, content_id)

    # 删除子内容后回退祖先状态（如删除 EPISODE → 父级 SEASON_SERIES 回退到 InProgress，
    # ApplicationReview 节点变为 Pending，需重新提交审核）
    if parent_id:
        from app.internal.cms_biz_orchestration.services.workflow_service import (
            rollback_ancestors_after_child_change,
            complete_process_and_update_status,
        )
        child_type_label = {"EPISODE": "单集", "SERIES": "单季", "SEASON_SERIES": "单季", "SCHEDULE": "节目单"}.get(
            content_type, content_type
        )
        await rollback_ancestors_after_child_change(
            db,
            start_parent_id=parent_id,
            edited_by=current_user.username,
            edit_info=f"删除{child_type_label}「{content_title}」",
        )
        # 在父级 Processes 中记录删除操作（Episodes/Season Series/Physical Channel 节点）
        # node_code=InjectSubContent 会在 list_processes 中按父级 content_type 显示对应名称
        parent_detail = await content_service.get_content(db, parent_id)
        if parent_detail and parent_detail.content.content_type in ("SERIES", "SEASON_SERIES", "SEASON", "CHANNEL"):
            await complete_process_and_update_status(
                db,
                content_id=parent_id,
                content_type=parent_detail.content.content_type,
                process_name="InjectSubContent",
                processed_by=current_user.username,
                info=f"删除{child_type_label}「{content_title}」",
                skip_status_update=True,
            )

    if content_type == ContentType.EPISODE.value:
        # 单集删除日志归属父级（content_id=parent_id），与 EPISODE_INJECT / SEASON_SERIES_REMOVE
        # 口径一致，保证父级（SERIES 普通连续剧 / SEASON_SERIES 单季）详情页 Activity Log 可见；
        # entity_id 仍指向被删单集自身，保留实体定位能力；无父级的顶层单集兜底挂自身
        await write_log(
            db,
            user_id=current_user.id,
            user_name=current_user.username,
            operation_type=OperationType.EPISODE_REMOVE,
            operation_object_code="OBJ_CONTENT", operation_object_params={"name": content_title},
            operation_content_code="log.episode.remove", operation_content_params={"title": content_title or ""},
            content_id=parent_id if parent_id else content_id,
            entity_type="content",
            entity_id=content_id,
            previous_value=prev_val,
            updated_value_json=raw_val,
            ip_address=_get_ip(request),
            result="success",
        )
    elif content_type == ContentType.SEASON_SERIES.value and parent_id:
        # 季详情页 Season Series 弹窗删除单季系列：日志归属父级（content_id=parent_id），
        # 与 SEASON_SERIES_INJECT 对称，保证父级（SEASON）详情页 Activity Log 可见；
        # entity_id 仍指向被删 SEASON_SERIES 自身，保留实体定位能力
        await write_log(
            db,
            user_id=current_user.id,
            user_name=current_user.username,
            operation_type=OperationType.SEASON_SERIES_REMOVE,
            operation_object_code="OBJ_CONTENT", operation_object_params={"name": content_title},
            operation_content_code="log.seasonSeries.remove", operation_content_params={"title": content_title or ""},
            content_id=parent_id,
            entity_type="content",
            entity_id=content_id,
            previous_value=prev_val,
            updated_value_json=raw_val,
            ip_address=_get_ip(request),
            result="success",
        )
    else:
        await write_log(
            db,
            user_id=current_user.id,
            user_name=current_user.username,
            operation_type=OperationType.CONTENT_DELETE,
            operation_object_code="OBJ_CONTENT", operation_object_params={"name": content_title},
            operation_content_code="LOG_CONTENT_DELETE", operation_content_params={"title": content_title},
            content_id=content_id,
            entity_type="content",
            entity_id=content_id,
            previous_value=prev_val,
            updated_value_json=raw_val,
            ip_address=_get_ip(request),
            result="success",
        )
    await db.commit()
    return {"success": True}


@router.post("/batch-delete")
async def batch_delete_contents_api(
    body: BatchDeleteRequest,
    request: Request,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    contents = (await db.execute(select(Content).where(Content.id.in_(body.ids)))).scalars().all()
    content_names = ", ".join([c.title for c in contents]) if contents else str(body.ids)
    prev_data = [orm_to_dict(c, "content") for c in contents]
    prev_val, _, raw_val = await prepare_log_values(db, "content", prev_data, None)
    deleted = await content_service.batch_delete_contents(db, body.ids)

    # 批量删除子内容后回退祖先状态（去重父级，避免同一父级重复回退）
    from app.internal.cms_biz_orchestration.services.workflow_service import rollback_ancestors_after_child_change
    for parent_id in {c.parent_id for c in contents if c.parent_id}:
        await rollback_ancestors_after_child_change(
            db,
            start_parent_id=parent_id,
            edited_by=current_user.username,
            edit_info=f"批量删除内容「{content_names}」",
        )

    await write_log(
        db,
        user_id=current_user.id,
        user_name=current_user.username,
        operation_type=OperationType.CONTENT_BATCH_DELETE,
        operation_object_code="OBJ_CONTENT", operation_object_params={"name": content_names},
        operation_content_code="LOG_CONTENT_BATCH_DELETE", operation_content_params={"names": content_names},
        ip_address=_get_ip(request),
        result="success",
        entity_type="content",
        entity_id=body.ids[0] if body.ids else None,
        previous_value=prev_val,
        updated_value=None,
        updated_value_json=raw_val,
    )
    await db.commit()
    return {"deleted": deleted}


@router.get("/{content_id}/licenses", response_model=list[ContentLicenseRef])
async def get_content_licenses(
    content_id: int,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(get_current_user),
):
    return await content_service.get_content_licenses(db, content_id)


@router.get("/{content_id}/adjacent", response_model=AdjacentContentResponse)
async def get_adjacent_content(
    content_id: int,
    content_types: str | None = Query(None, description="内容类型过滤，多个用逗号分隔，如：MOVIE,EPISODE"),
    is_archived: bool | None = Query(None, description="是否只查询归档内容"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """查询上一条/下一条内容 ID（用于详情页记录导航）。

    Args:
        content_id: 当前内容 ID
        content_types: 可选，内容类型过滤，如 "MOVIE,EPISODE"
        is_archived: 可选，是否只查询归档内容
    """
    # 解析 content_types 字符串为列表
    content_type_list = None
    if content_types:
        content_type_list = [t.strip() for t in content_types.split(",") if t.strip()]

    return await content_service.get_adjacent_content(
        db, content_id, current_user,
        content_types=content_type_list, is_archived=is_archived
    )


# ---------- Custom Field Values ----------

@router.get("/{content_id}/field-values", response_model=list[EntityFieldValueItem])
async def get_content_field_values(
    content_id: int,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(get_current_user),
):
    return await get_field_values(db, "Content", content_id)


@router.put("/{content_id}/field-values", response_model=list[EntityFieldValueItem])
async def save_content_field_values(
    content_id: int,
    body: EntityFieldValuesPayload,
    request: Request,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    from app.internal.cms_biz_package.models.enums import ContentType
    # 自定义字段编辑需要写操作日志的内容类型（频道/节目单/VOD）
    LOG_FIELD_TYPES = (
        ContentType.CHANNEL.value, ContentType.SCHEDULE.value,
        ContentType.MOVIE.value, ContentType.EPISODE.value,
        ContentType.SERIES.value, ContentType.SEASON_SERIES.value, ContentType.SEASON.value,
    )
    content_obj = (await db.execute(select(Content.content_type).where(Content.id == content_id))).scalar_one_or_none()
    # 已发布/已下架内容编辑后回滚状态并新建提交审核记录（与 Metadata 编辑行为一致）
    if content_obj:
        from app.internal.cms_biz_orchestration.services.workflow_service import rollback_after_published_edit
        await rollback_after_published_edit(
            db,
            content_id=content_id,
            content_type=content_obj,
            edited_by=current_user.username,
            edit_info="编辑自定义字段",
        )
    if content_obj in LOG_FIELD_TYPES:
        import json
        from app.internal.cms_biz_metada.models.basic import CustomField, CustomFieldOption
        from app.internal.cms_biz_metada.services.entity_data_service import get_field_values
        old_fv = await get_field_values(db, "Content", content_id)
        old_fv_map = {r.custom_field_id: r.value for r in old_fv}
    result = await save_field_values(db, "Content", content_id, body)
    if content_obj in LOG_FIELD_TYPES:
        field_ids = [v.custom_field_id for v in body.values if v.value is not None and v.value != ""]
        if field_ids:
            cf_rows = (await db.execute(select(CustomField).where(CustomField.id.in_(field_ids), CustomField.is_deleted.is_(False)))).scalars().all()
            field_info = {cf.id: cf for cf in cf_rows}
            dropdown_ids = [cf.id for cf in cf_rows if cf.field_type in ('DropList', 'DropList_multiple', 'multi_select')]
            option_map: dict[int, dict[str, str]] = {}
            if dropdown_ids:
                opt_rows = (await db.execute(select(CustomFieldOption).where(CustomFieldOption.custom_field_id.in_(dropdown_ids)))).scalars().all()
                for opt in opt_rows:
                    if opt.custom_field_id not in option_map:
                        option_map[opt.custom_field_id] = {}
                    names = opt.names or {}
                    display_name = resolve_option_display_name(names, opt.code)
                    option_map[opt.custom_field_id][opt.code] = display_name
            field_data = {}
            for v in body.values:
                if v.value is not None and v.value != "" and v.custom_field_id in field_info:
                    cf = field_info[v.custom_field_id]
                    if old_fv_map.get(v.custom_field_id) == v.value:
                        continue
                    raw_value = v.value
                    if cf.field_type in ('DropList', 'DropList_multiple', 'multi_select') and v.custom_field_id in option_map:
                        codes = raw_value.split(',')
                        translated = [option_map[v.custom_field_id].get(c.strip(), c.strip()) for c in codes]
                        field_data[cf.field_name] = ','.join(translated)
                    else:
                        field_data[cf.field_name] = raw_value
            if field_data:
                upd_val = json.dumps(field_data, ensure_ascii=False, default=str)
                if content_obj == ContentType.CHANNEL.value:
                    op_type = OperationType.CHANNEL_FIELD_UPDATE
                elif content_obj == ContentType.SCHEDULE.value:
                    op_type = OperationType.SCHEDULE_FIELD_UPDATE
                else:
                    op_type = OperationType.VOD_FIELD_UPDATE
                await write_log(
                    db,
                    user_id=current_user.id,
                    user_name=current_user.username,
                    operation_type=op_type,
                    operation_object_code="log.field.edit",
                    operation_content_code="log.field.edit",
                    content_id=content_id,
                    entity_type="content",
                    entity_id=content_id,
                    previous_value=None,
                    updated_value=upd_val,
                    # 原始数据快照（自定义字段变更映射）
                    updated_value_json=upd_val,
                    ip_address=_get_ip(request),
                    result="success",
                )
    await db.commit()
    return result


# ---------- Import Templates ----------

@router.get("/template/{content_type}")
async def get_import_template(
    content_type: str,
    _: User = Depends(get_current_user),
):
    """
    下载导入模板（EPISODE/SERIES）

    Args:
        content_type: 内容类型，支持 'EPISODE' 或 'SERIES'
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    if content_type not in ("EPISODE", "SERIES", "SEASON_SERIES"):
        from app.common.core.exceptions import BusinessException, ErrorCode
        from app.common.core.i18n import get_msg
        raise BusinessException(ErrorCode.VALIDATION_ERROR, get_msg("CONTENT_TYPE_INVALID"))

    wb = Workbook()
    ws = wb.active

    if content_type == "EPISODE":
        ws.title = "Episodes"
        headers = ["Episode Name", "Sequence", "Assignee"]
        filename = "Episode_Import_Template.xlsx"
        example_data = [
            ["Episode 1", 1, "admin"],
            ["Episode 2", 2, "admin"],
        ]
    else:  # SERIES
        ws.title = "Series"
        headers = ["Series Name", "Series Ordinal", "Assignee"]
        filename = "Series_Import_Template.xlsx"
        example_data = [
            ["Series 1", 1, "admin"],
            ["Series 2", 2, "admin"],
        ]

    # 表头样式
    header_fill = PatternFill(start_color="1677FF", end_color="1677FF", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    # 写入表头
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 写入示例数据
    for row_idx, row_data in enumerate(example_data, 2):
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    # 设置列宽
    col_widths = [30, 15, 20]
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    # 保存到内存
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.post("/parse-excel")
async def parse_excel_file(
    content_type: str,
    file: UploadFile,
    _: User = Depends(get_current_user),
):
    """
    解析 Excel 导入文件（EPISODE/SERIES）

    Args:
        content_type: 内容类型，支持 'EPISODE' 或 'SERIES'
        file: Excel 文件

    Returns:
        解析后的数据列表
    """
    from openpyxl import load_workbook

    if content_type not in ("EPISODE", "SERIES", "SEASON_SERIES"):
        from app.common.core.exceptions import BusinessException, ErrorCode
        from app.common.core.i18n import get_msg
        raise BusinessException(ErrorCode.VALIDATION_ERROR, get_msg("CONTENT_TYPE_INVALID"))

    # 读取文件内容
    contents = await file.read()
    buf = io.BytesIO(contents)

    # 加载工作簿（校验文件格式）
    try:
        wb = load_workbook(buf, data_only=True)
    except Exception:
        from app.common.core.exceptions import BusinessException, ErrorCode
        from app.common.core.i18n import get_msg
        raise BusinessException(ErrorCode.VALIDATION_ERROR, get_msg("INVALID_EXCEL_FORMAT"))
    ws = wb.active

    # 读取所有行
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(row)

    # 跳过表头，解析数据行
    result = []
    for i, row in enumerate(rows[1:], start=2):  # 从第2行开始（跳过表头）
        if not row or not row[0]:  # 跳过空行
            continue

        title = str(row[0]).strip() if row[0] else ""
        ordinal = int(row[1]) if row[1] and str(row[1]).isdigit() else None
        assignee = str(row[2]).strip() if row[2] else None

        if not title:  # 跳过标题为空的行
            continue

        item = {
            "row": i,
            "title": title,
            "assignee": assignee,
        }

        if content_type == "EPISODE":
            item["sequence"] = ordinal
        else:  # SERIES
            item["series_ordinal"] = ordinal

        result.append(item)

    return {"items": result}


@router.get("/{content_id}/node-status")
async def get_node_status(
    content_id: int,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(get_current_user),
):
    """基于实际数据关系判断各流程节点的完成状态"""
    return await content_service.get_node_completion_status(db, content_id)


# ---------- I18n Values ----------

@router.get("/{content_id}/i18n", response_model=list[EntityI18nItem])
async def get_content_i18n(
    content_id: int,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(get_current_user),
):
    return await get_i18n_values(db, "Content", content_id)


@router.put("/{content_id}/i18n", response_model=list[EntityI18nItem])
async def save_content_i18n(
    content_id: int,
    body: EntityI18nPayload,
    request: Request,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    from app.internal.cms_biz_package.models.enums import ContentType
    content_obj = (await db.execute(select(Content.content_type).where(Content.id == content_id))).scalar_one_or_none()
    # 已发布/已下架内容编辑后回滚状态并新建提交审核记录（与 Metadata 编辑行为一致）
    if content_obj:
        from app.internal.cms_biz_orchestration.services.workflow_service import rollback_after_published_edit
        await rollback_after_published_edit(
            db,
            content_id=content_id,
            content_type=content_obj,
            edited_by=current_user.username,
            edit_info="编辑多语言信息",
        )
    # 多语言字段编辑需要写操作日志的内容类型（频道/节目单/VOD）
    LOG_I18N_TYPES = (
        ContentType.CHANNEL.value, ContentType.SCHEDULE.value,
        ContentType.MOVIE.value, ContentType.EPISODE.value,
        ContentType.SERIES.value, ContentType.SEASON_SERIES.value, ContentType.SEASON.value,
    )
    if content_obj in LOG_I18N_TYPES:
        import json
        from app.internal.cms_biz_metada.services.entity_data_service import get_i18n_values
        from app.internal.cms_biz_metada.models.basic import CustomField, CustomFieldOption
        old_i18n = await get_i18n_values(db, "Content", content_id)
        old_map = {r.field_name: r.value for r in old_i18n if r.language == body.language}
        changed_fields = {}
        for k, v in body.fields.items():
            if v is not None and v != "" and old_map.get(k) != v:
                changed_fields[k] = v
        # 自定义字段变更单独写「自定义字段编辑」日志，其余多语言字段变更才写「多语言编辑」日志
        cf_changed = {k: v for k, v in changed_fields.items() if k.startswith('cf_')}
        cf_info: dict[str, CustomField] = {}
        option_map: dict[str, dict[str, str]] = {}
        if cf_changed:
            cf_rows = (await db.execute(select(CustomField).where(CustomField.field_code.in_(list(cf_changed.keys())), CustomField.is_deleted.is_(False)))).scalars().all()
            dropdown_ids = []
            for cf in cf_rows:
                cf_info[cf.field_code] = cf
                if cf.field_type in ('DropList', 'DropList_multiple', 'multi_select'):
                    dropdown_ids.append(cf.id)
            if dropdown_ids:
                opt_rows = (await db.execute(select(CustomFieldOption).where(CustomFieldOption.custom_field_id.in_(dropdown_ids)))).scalars().all()
                id_to_code = {cf.id: cf.field_code for cf in cf_rows if cf.id in dropdown_ids}
                for opt in opt_rows:
                    fc = id_to_code.get(opt.custom_field_id)
                    if fc and fc not in option_map:
                        option_map[fc] = {}
                    if fc:
                        names = opt.names or {}
                        # 选项名称与前端弹窗保持一致：按多语言编辑页签语言解析（bug 32055/32420）
                        display_name = resolve_option_display_name(names, opt.code, language=body.language)
                        option_map[fc][opt.code] = display_name
        cf_translated: dict[str, str] = {}
        for k, v in cf_changed.items():
            if k in cf_info:
                cf = cf_info[k]
                if cf.field_type in ('DropList', 'DropList_multiple', 'multi_select') and cf.field_code in option_map:
                    codes = v.split(',')
                    translated = [option_map[cf.field_code].get(c.strip(), c.strip()) for c in codes]
                    cf_translated[cf.field_name] = ','.join(translated)
                else:
                    cf_translated[cf.field_name] = v
            else:
                cf_translated[k] = v
        if cf_translated:
            cf_upd_val = json.dumps(cf_translated, ensure_ascii=False, default=str)
            if content_obj == ContentType.CHANNEL.value:
                cf_op_type = OperationType.CHANNEL_FIELD_UPDATE
            elif content_obj == ContentType.SCHEDULE.value:
                cf_op_type = OperationType.SCHEDULE_FIELD_UPDATE
            else:
                cf_op_type = OperationType.VOD_FIELD_UPDATE
            await write_log(
                db,
                user_id=current_user.id,
                user_name=current_user.username,
                operation_type=cf_op_type,
                operation_object_code="log.field.edit",
                # 带语言标识，显示"自定义字段编辑（en）"，区分各语言页签的变更（bug 32055）
                operation_content_code="log.field.edit.lang", operation_content_params={"lang_code": body.language},
                content_id=content_id,
                entity_type="content",
                entity_id=content_id,
                previous_value=None,
                updated_value=cf_upd_val,
                # 原始数据快照（自定义字段多语言变更映射）
                updated_value_json=cf_upd_val,
                ip_address=_get_ip(request),
                result="success",
            )
        translated_fields = {k: v for k, v in changed_fields.items() if not k.startswith('cf_')}
        fk_fields = {'tag_ids': ('tag', '标签'), 'genre_ids': ('genre', '题材')}
        for fk_key, (table, label) in fk_fields.items():
            if fk_key in translated_fields:
                from app.database import Base
                ids_str = translated_fields.pop(fk_key)
                ids = [int(x.strip()) for x in ids_str.split(',') if x.strip()]
                if ids:
                    model_cls = None
                    for mapper in Base.registry.mappers:
                        cls = mapper.class_
                        if getattr(cls, "__tablename__", None) == table:
                            model_cls = cls
                            break
                    if model_cls:
                        rows = (await db.execute(select(model_cls).where(model_cls.id.in_(ids), model_cls.language == body.language, model_cls.is_deleted.is_(False)))).scalars().all()
                        name_map = {r.id: r.name for r in rows}
                        names = [name_map.get(i, str(i)) for i in ids]
                        translated_fields[label] = ','.join(names)
                    else:
                        translated_fields[label] = ids_str
        if translated_fields:
            upd_val = json.dumps(translated_fields, ensure_ascii=False, default=str)
            if content_obj == ContentType.CHANNEL.value:
                op_type = OperationType.CHANNEL_I18N_UPDATE
            elif content_obj == ContentType.SCHEDULE.value:
                op_type = OperationType.SCHEDULE_I18N_UPDATE
            else:
                op_type = OperationType.VOD_I18N_UPDATE
            await write_log(
                db,
                user_id=current_user.id,
                user_name=current_user.username,
                operation_type=op_type,
                operation_object_code="log.i18n.edit",
                operation_content_code="log.i18n.edit.lang", operation_content_params={"lang_code": body.language},
                content_id=content_id,
                entity_type="content",
                entity_id=content_id,
                previous_value=None,
                updated_value=upd_val,
                # 原始数据快照（多语言字段变更映射）
                updated_value_json=upd_val,
                ip_address=_get_ip(request),
                result="success",
            )
        # 多语言变更的流程记录由元数据弹窗最后一步的 create/update_xxx_metadata
        # 统一写入（同一保存动作只记一条，此处不再补写，
        # 避免中间步骤写入 Pending 状态的重复记录）
    result = await save_i18n_values(db, "Content", content_id, body)
    await db.commit()
    return result
